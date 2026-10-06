"""Lưu các mẫu đã được CTV chấp nhận vào kho riêng da_duyet/ (tách khỏi output/, output_v2/ – chạy lại pipeline không ghi đè).

Với mỗi dòng được chấp nhận trong file Excel CTV, script kiểm tra:
  - mẫu ở trạng thái ĐẠT và có JSON chốt (json_vi/final, json_ko/final) trong thư mục nguồn;
  - bản tiếng Hàn hiện tại KHỚP với nội dung CTV đã xem (cột "Giải thích KOR JSON mới");
  - CTV không sửa ở cột "KO — CTV SỬA TẠI ĐÂY" (có sửa → chưa lưu, báo lại để xử lý riêng).
Đạt hết mới chép vào kho và ghi/cập nhật da_duyet/danh_sach.csv. Bản sao file Excel CTV lưu ở da_duyet/ctv_phan_hoi/.

Cách chạy (từ thư mục dạng bài):
  python scripts/s10_luu_da_duyet.py --file <excel> [--file <excel> …] [--nguon output] [--chap-nhan "Đạt"] [--tat-ca-dong] [--dry-run]
    --nguon        thư mục output chứa JSON (mặc định: "output" trong config.json)
    --chap-nhan    giá trị cột "Kết quả" được coi là chấp nhận (mặc định "Đạt"; nhiều giá trị cách nhau dấu phẩy)
    --tat-ca-dong  coi MỌI dòng của file là chấp nhận (file CTV đã duyệt nhưng không điền cột Kết quả)
    --ngay-duyet   YYYY-MM-DD (mặc định: lấy từ đầu tên file Excel)
    --ghi-de       cho phép thay mẫu đã có trong kho bằng bản khác
    --bo-mau       sample_id không lưu dù được chấm chấp nhận, VD "660_1" (CTV chấm Đạt nhưng vẫn góp ý)
"""
import argparse
import datetime
import hashlib
import json
import os
import re
import shutil
import sys

import _dang as D
from pipeline_chung.common import load_json, read_csv, write_csv

FIELDS = ["sample_id", "question_id", "cau_con", "cap_do", "ngay_duyet", "file_ctv", "ket_qua_ctv", "ghi_chu_ctv",
          "goi_y_han_han", "cach_chia_dang", "nguon_output", "hash_goc", "sha256_json_vi", "sha256_json_ko",
          "model_gd1", "prompt_gd1", "model_gd2", "prompt_gd2", "ngay_luu"]
COL_KO, COL_SUA, COL_KQ, COL_GC = "Giải thích KOR JSON mới", "KO — CTV SỬA TẠI ĐÂY", "Kết quả", "Ghi chú CTV"


def norm_text(s):
    """So nội dung như CTV nhìn thấy: bỏ ký tự điều khiển (đánh dấu gạch chân/in đậm), dấu ｜ của furigana, gộp khoảng trắng."""
    s = re.sub(r"[\x00-\x09\x0b-\x1f｜]", "", s or "")
    return re.sub(r"\s+", "", s)          # bỏ hẳn khoảng trắng: Excel (rich text) có thể nuốt khoảng trắng sau 】 / đầu dòng


def sha256(p):
    with open(p, "rb") as f:
        return hashlib.sha256(f.read()).hexdigest()


def han_han(note):
    n = (note or "").lower()
    if "hoặc không" in n:
        return "tuy_chon"
    if "không cần" in n:
        return "khong_can"
    if "한자음" in n or "hán-hàn" in n or "hán hàn" in n:
        return "co"
    return ""


def read_excel(p):
    from openpyxl import load_workbook
    wb = load_workbook(p, data_only=True)
    ws = next((w for w in wb.worksheets if w.title.upper().startswith("03")), wb.worksheets[1])
    rows = list(ws.iter_rows(values_only=True))
    h = rows[0]
    return [dict(zip(h, r)) for r in rows[1:] if r and r[0] is not None]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--file", action="append", required=True)
    ap.add_argument("--nguon")
    ap.add_argument("--chap-nhan", default="Đạt")
    ap.add_argument("--tat-ca-dong", action="store_true")
    ap.add_argument("--ngay-duyet")
    ap.add_argument("--ghi-de", action="store_true")
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--bo-mau", default="", help="sample_id không lưu dù CTV chấm chấp nhận (cách nhau dấu phẩy)")
    args = ap.parse_args()
    cfg = D.cfg()
    nguon = args.nguon or cfg["output"]
    os.environ["PIPELINE_OUTPUT"] = nguon
    st = {r["sample_id"]: r for r in read_csv(D.OUT("reports", "trang_thai.csv"))}
    norm_cd = "v2" if (nguon == cfg["output"] and cfg.get("cach_chia_dang") == "v2") else "v1"
    chap_nhan = {x.strip() for x in args.chap_nhan.split(",") if x.strip()}
    kho = D.da_duyet_rows()
    today = datetime.date.today().isoformat()
    them, bo = [], []

    for fx in args.file:
        fx = os.path.abspath(fx)
        name = os.path.basename(fx)
        m = re.match(r"(\d{4}-\d{2}-\d{2})", name)
        ngay = args.ngay_duyet or (m.group(1) if m else today)
        for x in read_excel(fx):
            sid = f"{x['question_id']}_{x['câu con']}"
            kq = (x.get(COL_KQ) or "").strip() if isinstance(x.get(COL_KQ), str) else x.get(COL_KQ)
            if not args.tat_ca_dong and kq not in chap_nhan:
                continue
            if sid in {b.strip() for b in args.bo_mau.split(",") if b.strip()}:
                bo.append((sid, name, "loại theo --bo-mau (CTV chấm đạt nhưng có góp ý cần xử lý)"))
                continue
            fv, fk = D.OUT("json_vi", "final", sid + ".json"), D.OUT("json_ko", "final", sid + ".json")
            s = st.get(sid) or {}
            ly = None
            if s.get("trang_thai") != "ĐẠT":
                ly = f"trạng thái {s.get('trang_thai') or 'không có'} (cần ĐẠT)"
            elif not (os.path.exists(fv) and os.path.exists(fk)):
                ly = "thiếu JSON chốt"
            elif norm_text(D.render(load_json(fk), cfg["ngon_ngu_dich"])) != norm_text(x.get(COL_KO)):
                ly = "bản KO hiện tại KHÁC nội dung CTV đã duyệt (đã chạy lại sau khi xuất Excel?)"
            elif x.get(COL_SUA) is not None and norm_text(x.get(COL_SUA)) != norm_text(x.get(COL_KO)):
                ly = "CTV có sửa ở cột KO — CTV SỬA TẠI ĐÂY → cần áp dụng chỗ sửa trước khi lưu"
            if not ly and sid in kho and kho[sid].get("sha256_json_ko") != sha256(fk) and not args.ghi_de:
                ly = "đã có trong kho với bản khác (thêm --ghi-de nếu muốn thay)"
            if ly:
                bo.append((sid, name, ly))
                continue
            row = {"sample_id": sid, "question_id": x["question_id"], "cau_con": x["câu con"], "cap_do": s.get("cap_do"),
                   "ngay_duyet": ngay, "file_ctv": name, "ket_qua_ctv": kq or "(chấp nhận – file không điền Kết quả)",
                   "ghi_chu_ctv": (x.get(COL_GC) or "").strip(), "goi_y_han_han": han_han(x.get(COL_GC)),
                   "cach_chia_dang": norm_cd, "nguon_output": nguon, "hash_goc": s.get("hash_goc"),
                   "sha256_json_vi": sha256(fv), "sha256_json_ko": sha256(fk),
                   "model_gd1": s.get("model_gd1"), "prompt_gd1": s.get("prompt_gd1"),
                   "model_gd2": s.get("model_gd2"), "prompt_gd2": s.get("prompt_gd2"), "ngay_luu": today}
            if sid in kho and kho[sid].get("sha256_json_ko") == row["sha256_json_ko"]:
                bo.append((sid, name, "đã có trong kho (giống hệt) – bỏ qua"))
                continue
            them.append(row)
            if not args.dry_run:
                for sub, src in (("json_vi", fv), ("json_ko", fk)):
                    os.makedirs(D.da_duyet_path(sub), exist_ok=True)
                    shutil.copy2(src, D.da_duyet_path(sub, sid + ".json"))
                os.makedirs(D.da_duyet_path("ly_do"), exist_ok=True)
                for gd, src in (("gd1", D.OUT("json_vi", "checked", sid + ".ly_do.json")),
                                ("gd2", D.OUT("json_ko", "checked", sid + ".ly_do.json"))):
                    if os.path.exists(src):
                        shutil.copy2(src, D.da_duyet_path("ly_do", f"{sid}.{gd}.ly_do.json"))
                kho[sid] = row
        if not args.dry_run:
            os.makedirs(D.da_duyet_path("ctv_phan_hoi"), exist_ok=True)
            dst = D.da_duyet_path("ctv_phan_hoi", name)
            if not os.path.exists(dst):
                shutil.copy2(fx, dst)

    if not args.dry_run and them:
        write_csv(sorted(kho.values(), key=lambda r: (r["cap_do"] or "", int(r["question_id"]), int(r["cau_con"]))),
                  D.da_duyet_path("danh_sach.csv"), FIELDS)
    if not args.dry_run and them:       # đánh dấu trong thư mục output đang dùng (config) để các bước sau bỏ qua
        os.environ["PIPELINE_OUTPUT"] = cfg["output"]
        if os.path.exists(D.OUT("reports", "trang_thai.csv")):
            from pipeline_chung.trang_thai import Status
            st2 = Status(D.OUT("reports", "trang_thai.csv"))
            n_dd, _ = D.danh_dau_da_duyet(st2)
            st2.save()
            print(f"Đánh dấu ĐÃ_DUYỆT_CTV trong {cfg['output']}/: {n_dd} câu")
    print(f"{'(dry-run) ' if args.dry_run else ''}Thêm vào kho: {len(them)} câu · bỏ qua: {len(bo)} · kho hiện có: {len(kho)} câu")
    for sid, name, ly in bo:
        print(f"  – {sid} ({name}): {ly}")
    if them:
        import collections
        print("  theo cấp độ (mới thêm):", dict(sorted(collections.Counter(r["cap_do"] for r in them).items())))
    print(f"→ {D.da_duyet_path()}")


if __name__ == "__main__":
    sys.exit(main())
