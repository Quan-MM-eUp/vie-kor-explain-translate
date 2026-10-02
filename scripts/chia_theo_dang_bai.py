"""Chia question-explain-INTEGRATED.csv thành 21 file CSV, mỗi file một dạng bài.

Danh sách 21 dạng lấy theo sheet "Format giải thích" của file
"Migii JLPT_Nâng cấp giải thích.xlsx". Tên file đầu ra dùng cùng số thứ tự và tên
với các ví dụ JSON trong explanation_schema/examples/ (01_tu_vung_cach_doc_kanji …).

Cách chạy (từ thư mục Vie-Kor):
  python scripts/chia_theo_dang_bai.py                 # ghi vào dang_bai/<dạng>/data/
  python scripts/chia_theo_dang_bai.py --force         # ghi đè nếu đã có file
  python scripts/chia_theo_dang_bai.py --input <csv> --output <thư mục gốc>

Đầu ra (mặc định thư mục gốc là dang_bai/):
  dang_bai/01_tu_vung_cach_doc_kanji/data/01_tu_vung_cach_doc_kanji.csv
  …
  dang_bai/21_nghe_hieu_tong_hop/data/21_nghe_hieu_tong_hop.csv
  dang_bai/00_khong_ro_dang_bai.csv   (chỉ tạo nếu có dòng không khớp dạng nào)
  dang_bai/_tong_hop.csv              (số dòng theo dạng và cấp độ)

Ghi chú:
  - Giữ nguyên mọi cột và thứ tự dòng của file gốc. Chỉ bỏ ký tự NUL (\\0) – file gốc có vài ký tự này
    làm thư viện csv của Python báo lỗi.
  - File ghi bằng UTF-8 có BOM để mở bằng Excel không lỗi font.
  - Nếu file xlsx có mặt và cài openpyxl, script đối chiếu danh sách dạng bài với sheet
    "Format giải thích" và cảnh báo nếu lệch.
"""
import argparse
import collections
import csv
import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))  # thư mục Vie-Kor
csv.field_size_limit(10**9)

# (số, phần, tên trong sheet "Format giải thích", giá trị cột dang_bai trong CSV, tên file)
DANG_BAI = [
    (1,  "Từ vựng",   "Cách đọc kanji",                      "cách đọc kanji",         "01_tu_vung_cach_doc_kanji"),
    (2,  "Từ vựng",   "Thay đổi cách nói",                   "thay đổi cách nói",      "02_tu_vung_thay_doi_cach_noi"),
    (3,  "Từ vựng",   "Cách viết từ",                        "cách viết từ",           "03_tu_vung_cach_viet_tu"),
    (4,  "Từ vựng",   "Điền từ theo văn cảnh",               "điền từ theo văn cảnh",  "04_tu_vung_dien_tu_theo_van_canh"),
    (5,  "Từ vựng",   "Hình thành từ",                       "hình thành từ",          "05_tu_vung_hinh_thanh_tu"),
    (6,  "Từ vựng",   "Ứng dụng từ",                         "ứng dụng từ",            "06_tu_vung_ung_dung_tu"),
    (7,  "Ngữ pháp",  "Lựa chọn ngữ pháp",                   "lựa chọn ngữ pháp",      "07_ngu_phap_lua_chon_ngu_phap"),
    (8,  "Ngữ pháp",  "Lắp ghép câu",                        "lắp ghép câu",           "08_ngu_phap_lap_ghep_cau"),
    (9,  "Ngữ pháp",  "Ngữ pháp theo đoạn văn",              "ngữ pháp theo đoạn văn", "09_ngu_phap_theo_doan_van"),
    (10, "Đọc hiểu",  "Đoạn văn ngắn",                       "đoạn văn ngắn",          "10_doc_hieu_doan_van_ngan"),
    # CSV gọi là "đoạn văn vừa"
    (11, "Đọc hiểu",  "Đoạn văn trung bình",                 "đoạn văn vừa",           "11_doc_hieu_doan_van_trung_binh"),
    (12, "Đọc hiểu",  "Đoạn văn dài",                        "đoạn văn dài",           "12_doc_hieu_doan_van_dai"),
    (13, "Đọc hiểu",  "Đọc hiểu tổng hợp",                   "đọc hiểu tổng hợp",      "13_doc_hieu_tong_hop"),
    # CSV gọi là "đọc hiểu chủ đề" (chỉ có ở N1, N2 – phần 主張理解)
    (14, "Đọc hiểu",  "Đọc hiểu quan điểm (đoạn văn dài)",   "đọc hiểu chủ đề",        "14_doc_hieu_chu_de_quan_diem"),
    (15, "Đọc hiểu",  "Tìm thông tin",                       "tìm thông tin",          "15_doc_hieu_tim_thong_tin"),
    (16, "Nghe hiểu", "Nghe hiểu chủ đề",                    "nghe hiểu chủ đề",       "16_nghe_hieu_chu_de"),
    (17, "Nghe hiểu", "Nghe hiểu điểm chính",                "nghe hiểu điểm chính",   "17_nghe_hieu_diem_chinh"),
    (18, "Nghe hiểu", "Nghe hiểu khái quát",                 "nghe hiểu khái quát",    "18_nghe_hieu_khai_quat"),
    (19, "Nghe hiểu", "Nghe hiểu diễn đạt",                  "nghe hiểu diễn đạt",     "19_nghe_hieu_dien_dat"),
    (20, "Nghe hiểu", "Trả lời nhanh",                       "trả lời nhanh",          "20_nghe_hieu_tra_loi_nhanh"),
    (21, "Nghe hiểu", "Nghe hiểu tổng hợp",                  "nghe hiểu tổng hợp",     "21_nghe_hieu_tong_hop"),
]
KHONG_RO = "00_khong_ro_dang_bai"
LEVELS = ["N1", "N2", "N3", "N4", "N5"]


def norm(s):
    return " ".join((s or "").split()).lower()


def check_sheet(xlsx_path):
    """Đối chiếu tên dạng bài trong sheet với bảng DANG_BAI. Chỉ cảnh báo, không dừng."""
    try:
        import openpyxl
    except ImportError:
        print("(Bỏ qua đối chiếu sheet: chưa cài openpyxl)")
        return
    if not os.path.exists(xlsx_path):
        print(f"(Bỏ qua đối chiếu sheet: không thấy {os.path.basename(xlsx_path)})")
        return
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    if "Format giải thích" not in wb.sheetnames:
        print("CẢNH BÁO: file xlsx không có sheet 'Format giải thích'")
        return
    ws = wb["Format giải thích"]
    in_sheet = [norm(r[1]) for i, r in enumerate(ws.iter_rows(values_only=True)) if i > 0 and r[1]]
    expected = [norm(d[2]) for d in DANG_BAI]
    missing = [d for d in expected if d not in in_sheet]
    extra = [d for d in in_sheet if d not in expected]
    if missing or extra:
        print("CẢNH BÁO: danh sách dạng bài trong sheet khác với script")
        if missing:
            print("  - có trong script, không có trong sheet:", missing)
        if extra:
            print("  - có trong sheet, không có trong script:", extra)
    else:
        print(f"Đối chiếu sheet 'Format giải thích': khớp đủ {len(expected)} dạng bài.")


def main():
    ap = argparse.ArgumentParser(description="Chia file CSV lời giải thích theo 21 dạng bài")
    ap.add_argument("--input", default=os.path.join(ROOT, "question-explain-INTEGRATED.csv"))
    ap.add_argument("--output", default=os.path.join(ROOT, "dang_bai"), help="thư mục chứa 21 folder dạng bài")
    ap.add_argument("--xlsx", default=os.path.join(ROOT, "Migii JLPT_Nâng cấp giải thích.xlsx"))
    ap.add_argument("--force", action="store_true", help="ghi đè file đã có trong thư mục đầu ra")
    args = ap.parse_args()

    check_sheet(args.xlsx)

    def out_path(stem):
        """Dạng bài → <gốc>/<dạng>/data/<dạng>.csv; dòng không rõ dạng → <gốc>/00_khong_ro_dang_bai.csv."""
        if stem == KHONG_RO:
            return os.path.join(args.output, stem + ".csv")
        return os.path.join(args.output, stem, "data", stem + ".csv")

    existing = [out_path(d[4]) for d in DANG_BAI if os.path.exists(out_path(d[4]))]
    if existing and not args.force:
        sys.exit(f"Đã có {len(existing)} file đầu ra (vd. {existing[0]}). Dùng --force để ghi đè.")
    os.makedirs(args.output, exist_ok=True)

    lookup = {norm(d[3]): d for d in DANG_BAI}

    with open(args.input, encoding="utf-8-sig", newline="") as f:
        text = f.read()
    nul = text.count("\0")
    reader = csv.DictReader(text.replace("\0", "").splitlines(keepends=True))
    fields = reader.fieldnames
    groups = collections.defaultdict(list)
    unknown_values = collections.Counter()
    total = 0
    for row in reader:
        total += 1
        d = lookup.get(norm(row.get("dang_bai")))
        if d:
            groups[d[4]].append(row)
        else:
            groups[KHONG_RO].append(row)
            unknown_values[row.get("dang_bai", "")] += 1

    written = 0
    summary = []
    for num, phan, ten_sheet, ten_csv, stem in DANG_BAI + [(0, "", "Không rõ dạng bài", "", KHONG_RO)]:
        rows = groups.get(stem, [])
        if stem == KHONG_RO and not rows:
            continue
        path = out_path(stem)
        os.makedirs(os.path.dirname(path), exist_ok=True)
        with open(path, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=fields)
            w.writeheader()
            w.writerows(rows)
        written += len(rows)
        lv = collections.Counter(r["cap_do"] for r in rows)
        summary.append({
            "so": num, "phan": phan, "dang_bai_sheet": ten_sheet, "dang_bai_csv": ten_csv,
            "file": os.path.relpath(path, args.output).replace(os.sep, "/"),
            "so_dong": len(rows), **{lv_: lv.get(lv_, 0) for lv_ in LEVELS},
            "trong_giai_thich_vi": sum(1 for r in rows if not (r.get("giai_thich_vi") or "").strip()),
        })

    with open(os.path.join(args.output, "_tong_hop.csv"), "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(summary[0].keys()))
        w.writeheader()
        w.writerows(summary)

    print(f"Đọc {total} dòng (đã bỏ {nul} ký tự NUL) → ghi {written} dòng vào {len(summary)} file tại {args.output}")
    for s in summary:
        print(f"  {s['file']:75} {s['so_dong']:6}  " + " ".join(f"{l}:{s[l]}" for l in LEVELS)
              + (f"  | trống giai_thich_vi: {s['trong_giai_thich_vi']}" if s["trong_giai_thich_vi"] else ""))
    if unknown_values:
        print("Dòng không khớp dạng bài nào (giá trị dang_bai → số dòng):", dict(unknown_values))
    if written != total:
        sys.exit(f"LỖI: số dòng ghi ra ({written}) khác số dòng đọc vào ({total})")
    print("Kiểm tra: tổng số dòng khớp với file gốc.")


if __name__ == "__main__":
    main()
