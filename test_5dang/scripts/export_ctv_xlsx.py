"""Xuất bản dịch ra file Excel theo đúng format cộng tác viên (CTV) đã quen (tham khảo thư mục ../to_ctv).

Mỗi mẫu là 1 dòng, 17 cột giống file CTV: thông tin câu hỏi, Giải thích VI (tham chiếu),
Giải thích KO hiện tại (bản dịch mới), cột VÀNG "KO — CTV SỬA TẠI ĐÂY" (điền sẵn bản copy),
"Kết quả" (Đạt / Sửa nhỏ / Không đạt) và "Ghi chú CTV".

Lời giải được dựng lại từ JSON thành văn bản có nhãn như dữ liệu gốc:
  - Dạng cách đọc kanji: khối nhãn 【Câu hỏi】【Cách đọc câu】…【Tham khảo】 (nhãn giữ tiếng Việt như file CTV)
  - Các dạng khác: "Cách đọc: / Nghĩa: / ### / PHÂN TÍCH …" (VI) và "읽기: / 의미: / ### / 분석 …" (KO)
  - Phần gạch chân { } và các tiêu đề được IN ĐẬM thật trong Excel; dấu ⟪ ⟫ được bỏ.

Cách chạy (từ thư mục test_5dang):
  python scripts/export_ctv_xlsx.py --model opus --vi-base opus
  python scripts/export_ctv_xlsx.py --model sonnet            # vi-base mặc định: final
Kết quả: reports/ctv/<ngày>-CTV-test5dang_<model>.xlsx
"""
import argparse
import csv
import datetime
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import clean_html, load_config, load_json, load_raw, load_samples, p  # noqa: E402

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

B0, B1 = "\x01", "\x02"          # đánh dấu bắt đầu / kết thúc IN ĐẬM
U0, U1 = "\x03", "\x04"          # đánh dấu bắt đầu / kết thúc GẠCH CHÂN

# ---------------- nhãn & mẫu câu theo ngôn ngữ (sửa tại đây nếu CTV muốn đổi) ----------------
LABELS = {
    "vi": {"q": "Câu hỏi: ", "read": "Cách đọc: ", "mean": "Nghĩa: ", "analysis": "PHÂN TÍCH",
           "correct": "LỰA CHỌN ĐÚNG", "ref": "THÔNG TIN THAM KHẢO", "full": "Câu hoàn chỉnh: ",
           "meaning_tpl": 'Nghĩa là "{m}". ', "dict": "Thể từ điển: "},
    "ko": {"q": "문제: ", "read": "읽기: ", "mean": "의미: ", "analysis": "분석",
           "correct": "정답", "ref": "참고 정보", "full": "완성된 문장: ",
           "meaning_tpl": '"{m}"라는 의미입니다. ', "dict": "사전형: "},
}
KANJI_LABELS = ["Câu hỏi", "Cách đọc câu", "Nghĩa câu", "Từ", "Cách đọc từ", "Nghĩa từ",
                "Thể từ điển", "Đáp án đúng", "Tham khảo"]

HEADERS = ["question_id", "câu con", "Cấp độ", "Dạng bài", "Loại đề", "Free/Premium", "Nhóm review",
           "Điểm cần chú ý", "Câu hỏi", "Các lựa chọn", "Đáp án đúng (DB)", "Ngữ liệu (đoạn văn / transcript)",
           "Giải thích VI (tham chiếu)", "Giải thích KO hiện tại", "KO — CTV SỬA TẠI ĐÂY", "Kết quả", "Ghi chú CTV"]
WIDTHS = {"A": 10, "B": 6, "C": 7, "D": 16, "E": 14, "F": 9, "G": 18, "H": 26, "I": 30, "J": 20, "K": 8,
          "L": 40, "M": 55, "N": 55, "O": 55, "P": 11, "Q": 30}
HEAD_FILL, EDIT_HEAD_FILL = "305496", "BF8F00"
REF_FILL, EDIT_FILL = "EDEDED", "FFF2CC"


# ---------------- dựng văn bản từ JSON ----------------
def t(obj, lang):
    """Lấy văn bản theo ngôn ngữ từ trường cần dịch; đổi {…} thành in đậm, bỏ ⟪ ⟫."""
    if not obj:
        return ""
    s = obj.get(lang) or ""
    return fmt(s)


def fmt(s):
    s = (s or "").replace("⟪", "").replace("⟫", "")
    return s.replace("{", B0).replace("}", B1)


def bold(s):
    return f"{B0}{s}{B1}"


def opt_line(o, lang, L, with_meaning=True):
    head = f"{o['index']}. {o.get('option_ja', '')}"
    if o.get("reading") and o.get("reading") != o.get("option_ja"):
        head += f" ({o['reading']})"
    m = t(o.get("meaning"), lang)
    a = t(o.get("analysis"), lang)
    body = ""
    if with_meaning and m and m not in a:
        body += L["meaning_tpl"].format(m=m)
    body += a
    return f"{head}: {body}".rstrip(": ").rstrip()


def sentence_block(q, lang, L, q_label=None):
    lines = [(q_label or "") + fmt(q.get("ja", ""))]
    if q.get("reading"):
        lines.append(L["read"] + fmt(q["reading"]))
    if q.get("meaning"):
        lines.append(L["mean"] + t(q["meaning"], lang))
    return lines


def correct_block(d, lang, L, with_full=False):
    c = d.get("correct") or {}
    lines = [bold(L["correct"]), f"{c.get('index')}. {c.get('option_ja', '')}"]
    fs = c.get("full_sentence")
    if with_full and fs:
        lines.append(L["full"] + fmt(fs.get("ja", "")))
        if fs.get("meaning"):
            lines.append(L["mean"] + t(fs["meaning"], lang))
    return lines


def render(d, lang):
    L = LABELS[lang]
    sch = d.get("schema", "").split(".")[0]
    out = []

    if sch == "vocab_kanji_reading":
        w, q, c = d.get("word") or {}, d.get("question") or {}, d.get("correct") or {}
        dic = w.get("dictionary_form")
        vals = [fmt(q.get("ja", "")), fmt(q.get("reading") or ""), t(q.get("meaning"), lang),
                w.get("ja", ""), w.get("reading", ""), t(w.get("meaning"), lang),
                f"{dic['ja']} ({dic['reading']})" if dic else "", f"{c.get('index')}. {c.get('option_ja', '')}"]
        for lab, v in zip(KANJI_LABELS[:8], vals):
            out.append(f"{bold('【' + lab + '】')} {v}")
        for o in d.get("options", []):
            out.append(f"{bold('【Lựa chọn ' + str(o['index']) + ': ' + o.get('option_ja', '') + '】')} {t(o.get('analysis'), lang)}")
        out.append(bold("【Tham khảo】") + " ")
        out.append(t(d.get("reference"), lang))
        return "\n".join(out)

    if sch == "vocab_writing":
        out += sentence_block(d["question"], lang, L, L["q"]) + ["###"]
        out += correct_block(d, lang, L) + ["###", bold(L["analysis"])]
        for o in d.get("options", []):
            out.append(opt_line(o, lang, L))
        ref = d.get("reference") or {}
        if ref.get("text") or ref.get("examples"):
            out += ["###", bold(L["ref"])]
            if ref.get("text"):
                out.append(t(ref["text"], lang))
            for e in ref.get("examples", []):
                out.append(f"- {fmt(e.get('ja', ''))}" + (f"\n  {e['reading']}" if e.get("reading") else "") +
                           (f"\n  {t(e.get('meaning'), lang)}" if e.get("meaning") else ""))
        return "\n".join(out)

    if sch == "vocab_synonym":
        out += sentence_block(d["question"], lang, L) + ["###", bold(L["analysis"])]
        for o in d.get("options", []):
            out.append(opt_line(o, lang, L))
        out += ["###"] + correct_block(d, lang, L)
        ref = d.get("reference") or {}
        if ref.get("intro") or ref.get("terms"):
            out += ["###", bold(L["ref"])]
            if ref.get("intro"):
                out.append(t(ref["intro"], lang))
            for i, term in enumerate(ref.get("terms", []), 1):
                line = f"{i}. {term.get('ja', '')}" + (f" ({term['reading']})" if term.get("reading") else "")
                line += f": {t(term.get('meaning'), lang)}"
                if term.get("note"):
                    line += f" {t(term['note'], lang)}"
                out.append(line)
        return "\n".join(out)

    if sch == "vocab_context_fill":
        out += sentence_block(d["question"], lang, L) + ["###", bold(L["analysis"])]
        for o in d.get("options", []):
            out.append(opt_line(o, lang, L))
        out += ["###"] + correct_block(d, lang, L, with_full=True)
        return "\n".join(out)

    if sch == "vocab_word_formation":
        out += sentence_block(d["question"], lang, L) + ["###", bold(L["analysis"])]
        for o in d.get("options", []):
            head = f"{o['index']}. {o.get('option_ja', '')}" + (f" ({o['reading']})" if o.get("reading") else "")
            if o.get("gloss"):
                head += f" - {t(o['gloss'], lang)}"
            out.append(head)
            comp = o.get("compound") or {}
            line = "   " + fmt(comp.get("ja", ""))
            if comp.get("reading"):
                line += f" ({comp['reading']})"
            extra = " ".join(x for x in (t(comp.get("meaning"), lang), t(o.get("analysis"), lang)) if x)
            out.append(f"{line}: {extra}" if extra else line)
        out += ["###"] + correct_block(d, lang, L, with_full=True)
        ref = d.get("reference") or {}
        if ref.get("intro") or ref.get("terms"):
            out += ["###", bold(L["ref"])]
            if ref.get("intro"):
                out.append(t(ref["intro"], lang))
            for term in ref.get("terms", []):
                out.append(f"{term.get('ja', '')}" + (f" ({term['reading']})" if term.get("reading") else "") +
                           f": {t(term.get('meaning'), lang)}")
        return "\n".join(out)

    raise ValueError(f"Chưa hỗ trợ schema {sch} – bổ sung hàm render cho dạng bài này.")


# ---------------- chuyển chuỗi có đánh dấu thành rich text ----------------
def rich(s):
    if not any(m in s for m in (B0, U0)):
        return s
    parts, buf, b, u = [], "", False, False

    def flush():
        nonlocal buf
        if buf:
            parts.append(TextBlock(InlineFont(b=b or None, u="single" if u else None), buf) if (b or u) else buf)
            buf = ""
    for ch in s:
        if ch in (B0, B1, U0, U1):
            flush()
            b = True if ch == B0 else (False if ch == B1 else b)
            u = True if ch == U0 else (False if ch == U1 else u)
        else:
            buf += ch
    flush()
    return CellRichText(parts)


def question_text(html):
    """Câu hỏi gốc: giữ gạch chân (<u>) bằng định dạng thật."""
    s = re.sub(r"<u[^>]*>", U0, html or "")
    s = s.replace("</u>", U1)
    s = clean_html(s)
    return s


# ---------------- thông tin bổ sung từ CSV gốc ----------------
def csv_meta(cfg, keys):
    meta = {}
    with open(p(cfg["source_csv"]), encoding="utf-8-sig") as f:
        for row in csv.DictReader(line.replace("\0", "") for line in f):
            k = (row["question_id"], row["cau_con"])
            if k in keys:
                thi_thu = row.get("thi_thu_thuoc_de", "").strip()
                meta[k] = {"loai_de": "Thi thử JLPT" if thi_thu else "Luyện tập",
                           "free": (row.get("thi_thu_mo_cho") if thi_thu else row.get("luyen_tap_mo_cho")) or ""}
    return meta


def attention(sid, model, base):
    """Điểm cần chú ý: lỗi/cảnh báo tự động + ghi chú của Claude khi chuyển JSON và khi dịch."""
    items = []
    for path, key in ((p("reports", f"check_ko_{model}.json"), "check"), (p("reports", f"validate_vi_{base}.json"), "vi")):
        if os.path.exists(path):
            for r in load_json(path)["results"]:
                if r["sample_id"] != sid:
                    continue
                items += r.get("issues", [])
                if key == "check":
                    items += [f"{f['path']}: {e['loai']}" for f in r.get("fields", []) for e in f["errors"]
                              if e["muc_do"] != "nhẹ"]
    for path, tag in ((p("json_vi", base, "_notes.json"), "Bản gốc VI"), (p("json_ko", model, "_notes.json"), "Khi dịch")):
        if os.path.exists(path):
            for n in load_json(path):
                if isinstance(n, dict) and n.get("sample_id") == sid:
                    items.append(f"[{tag}] {n.get('mo_ta', '')}")
    return "\n".join(f"- {x}" for x in items)


# ---------------- ghi Excel ----------------
def guide_sheet(ws, model, n):
    rows = [
        (f"TEST 5 DẠNG — KO do {model.upper()} dịch từ JSON", None),
        ("Bản dịch mới được tạo bằng quy trình: tách lời giải VI thành JSON → AI dịch phần tiếng Việt → dựng lại văn bản. "
         "Cần CTV đánh giá chất lượng trước khi áp dụng cho toàn bộ dữ liệu.", None),
        ("Phạm vi", "5 dạng: cách đọc kanji, thay đổi cách nói, cách viết từ, điền từ theo văn cảnh, hình thành từ · "
                    f"ngày tạo {datetime.date.today().isoformat()}"),
        ("Số câu", f"{n}"),
        (None, None),
        ("Nhiệm vụ", 'So sánh cột "Giải thích KO hiện tại" với "Giải thích VI (tham chiếu)". Chọn Kết quả; nếu cần sửa thì sửa '
                     'trực tiếp ở cột VÀNG "KO — CTV SỬA TẠI ĐÂY" (đã điền sẵn bản copy KO hiện tại).'),
        ("4 tiêu chí", "① Đủ ý như VI, không thêm/bớt  ② Tiếng Nhật (câu hỏi, lựa chọn, trích dẫn) giữ nguyên  "
                       "③ 정답 và lựa chọn khớp \"Đáp án đúng (DB)\"  ④ Tiếng Hàn tự nhiên, đúng thuật ngữ (의미, 읽기, 분석, 정답, 참고 정보)"),
        ("Kết quả", "Đạt = không cần sửa · Sửa nhỏ = đã sửa ở cột vàng · Không đạt = sai nghiêm trọng, ghi lý do ở \"Ghi chú CTV\""),
        ("Lưu ý", "KO phải bám VI. Nếu thấy chính VI sai (sai đáp án, dịch sai) → ghi ở Ghi chú, KHÔNG tự sửa KO theo ý mình. "
                  'Cột "Điểm cần chú ý" liệt kê chỗ kiểm tra tự động hoặc AI nghi có vấn đề — nên xem trước.'),
        (None, None),
        ("CÁCH ĐỌC / SỬA NỘI DUNG", None),
        ("Xuống dòng", "Xuống dòng thật trong ô. Muốn xuống dòng khi sửa: Alt + Enter (Mac: Option + Enter)."),
        ("In đậm / gạch chân", "Hiển thị bằng định dạng thật của Excel (phần gạch chân của đề được in đậm trong lời giải). "
                               "Muốn in đậm khi sửa: bôi đen chữ → Ctrl+B. KHÔNG gõ ký hiệu ** hay __."),
        ("Dạng cách đọc kanji", "Nội dung tách thành các khối có nhãn 【Câu hỏi】【Cách đọc câu】【Nghĩa câu】【Lựa chọn 1: …】【Tham khảo】… "
                                "Chỉ sửa phần chữ SAU nhãn. KHÔNG sửa, xoá hay đổi thứ tự các nhãn 【…】."),
        ("Cột được sửa", "Chỉ sửa ở các cột nền VÀNG (tiêu đề nâu). Các cột nền xám / trắng là dữ liệu tham chiếu — KHÔNG sửa, "
                         "không xoá, không đổi thứ tự cột hay đổi tên sheet."),
    ]
    for r in rows:
        ws.append(list(r))
    ws["A1"].font = Font(bold=True, size=14)
    ws["A11"].font = Font(bold=True)
    for r in range(3, ws.max_row + 1):
        ws.cell(r, 1).font = Font(bold=True)
        for c in (1, 2):
            ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:B2")
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 24, 100


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--model", required=True, help="thư mục bản dịch trong json_ko/ (vd. opus)")
    ap.add_argument("--vi-base", default="final", help="thư mục JSON tiếng Việt đã dùng để dịch (vd. opus, final)")
    ap.add_argument("--lang", default=None, help="mã ngôn ngữ đích (mặc định theo config.target_lang)")
    args = ap.parse_args()
    cfg = load_config()
    lang = args.lang or cfg["target_lang"]
    samples = load_samples()
    meta = csv_meta(cfg, {(s["question_id"], s["cau_con"]) for s in samples})

    wb = Workbook()
    guide_sheet(wb.active, args.model, len(samples))
    wb.active.title = "HƯỚNG DẪN"
    ws = wb.create_sheet(f"TEST 5 DANG {args.model.upper()} ({len(samples)})")
    ws.append(HEADERS)
    missing = []
    for s in samples:
        sid = s["sample_id"]
        ko_path = p("json_ko", args.model, sid + ".json")
        if not os.path.exists(ko_path):
            missing.append(sid)
            continue
        d = load_json(ko_path)
        raw = load_raw(sid)
        m = meta.get((s["question_id"], s["cau_con"]), {})
        opts = "\n".join(f"{i}. {clean_html(raw[f'lua_chon_{i}'])}" for i in range(1, 5) if raw.get(f"lua_chon_{i}"))
        corpus = clean_html(raw.get("doan_van", "") or raw.get("transcript_vi", ""))
        vi_txt, ko_txt = render(d, "vi"), render(d, lang)
        ws.append([raw["question_id"], raw["cau_con"], raw["cap_do"], raw["dang_bai"], m.get("loai_de", ""),
                   m.get("free", ""), f"KO do {args.model.upper()} dịch từ JSON (test 5 dạng)",
                   attention(sid, args.model, args.vi_base), rich(question_text(raw["cau_hoi"])), opts,
                   raw["dap_an_so"], corpus or None, rich(vi_txt), rich(ko_txt), rich(ko_txt), None, None])

    # định dạng giống file CTV
    for c in ws[1]:
        edit = c.column_letter in ("O", "P", "Q")
        c.fill = PatternFill("solid", fgColor=EDIT_HEAD_FILL if edit else HEAD_FILL)
        c.font = Font(bold=True, color="FFFFFF")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if c.column_letter in ("I", "J", "K", "L", "M", "N"):
                c.fill = PatternFill("solid", fgColor=REF_FILL)
            elif c.column_letter in ("O", "P", "Q"):
                c.fill = PatternFill("solid", fgColor=EDIT_FILL)
        ws.row_dimensions[row[0].row].height = 405
    ws.row_dimensions[1].height = 43.5
    for col, w in WIDTHS.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:Q{ws.max_row}"
    dv = DataValidation(type="list", formula1='"Đạt,Sửa nhỏ,Không đạt"', allow_blank=True)
    dv.add(f"P2:P{ws.max_row}")
    ws.add_data_validation(dv)

    os.makedirs(p("reports", "ctv"), exist_ok=True)
    out = p("reports", "ctv", f"{datetime.date.today().isoformat()}-CTV-test5dang_{args.model}.xlsx")
    wb.save(out)
    print(f"→ {out}  ({ws.max_row - 1} câu)")
    if missing:
        print("Thiếu bản dịch, bỏ qua:", missing)


if __name__ == "__main__":
    main()
