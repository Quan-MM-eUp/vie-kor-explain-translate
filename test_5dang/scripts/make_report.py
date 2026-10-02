"""Bước 6–7 – Tổng hợp kết quả ra Excel và tạo phiếu chấm mù cho biên tập viên.

Đọc (file nào chưa có thì bỏ qua): reports/validate_vi_<model>.json, reports/check_ko_<model>.json,
reports/judge_results.json, và (tùy chọn) phiếu biên tập viên đã chấm.

Cách chạy:
  python scripts/make_report.py
      → reports/bao_cao_so_sanh.xlsx      (tổng hợp cho bạn/Leader)
      → reports/phieu_bien_tap_vien.xlsx  (phiếu chấm mù A/B cho biên tập viên; chỉ tạo nếu chưa có)
  python scripts/make_report.py --editor reports/phieu_bien_tap_vien.xlsx
      → đưa điểm biên tập viên đã chấm vào báo cáo
"""
import argparse
import collections
import os
import random
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import load_config, load_json, load_samples, p, save_json  # noqa: E402

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:
    sys.exit("Cần cài openpyxl: pip install openpyxl")

HEAD = PatternFill("solid", fgColor="DDE7F3")


def read(path):
    return load_json(path) if os.path.exists(path) else None


def style(ws, widths):
    for c in ws[1]:
        c.font, c.fill = Font(bold=True), HEAD
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"


def editor_scores(path, lang):
    """Đọc phiếu biên tập viên: trả về {model: [điểm...]}, danh sách ghi chú.
    Đáp án A/B nằm riêng trong reports/_dap_an_phieu_bien_tap_vien.json (không gửi file này cho biên tập viên)."""
    wb = load_workbook(path)
    mapping = {(r["sample_id"], r["path"]): {"A": r["A"], "B": r["B"]}
               for r in load_json(p("reports", "_dap_an_phieu_bien_tap_vien.json"))}
    scores, notes = collections.defaultdict(list), []
    for r in wb["Cham_diem"].iter_rows(min_row=2, values_only=True):
        sid, path_, _vi, _a, _b, sa, sb, note = (list(r) + [None] * 8)[:8]
        m = mapping.get((sid, path_))
        if not m:
            continue
        for lab, sc in (("A", sa), ("B", sb)):
            if isinstance(sc, (int, float)):
                scores[m[lab]].append(sc)
        if note:
            notes.append((sid, path_, note))
    return scores, notes


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--editor", help="đường dẫn phiếu biên tập viên đã chấm")
    args = ap.parse_args()
    cfg = load_config()
    models, lang = cfg["models"], cfg["target_lang"]
    samples = load_samples()
    types = list(cfg["dang_bai"])

    val = {m: read(p("reports", f"validate_vi_{m}.json")) for m in models}
    chk = {m: read(p("reports", f"check_ko_{m}.json")) for m in models}
    judge = read(p("reports", "judge_results.json"))
    ed_scores, ed_notes = editor_scores(args.editor, lang) if args.editor else ({}, [])

    wb = Workbook()

    # --- Tổng hợp ---
    ws = wb.active
    ws.title = "Tong_hop"
    ws.append(["Tiêu chí", "Dạng bài"] + models)
    for d in ["(tất cả)"] + types:
        ids = {s["sample_id"] for s in samples if d == "(tất cả)" or s["dang_bai"] == d}

        def pct(data, key="status", ok=("ĐẠT", "CẢNH BÁO")):
            if not data:
                return "–"
            rs = [r for r in data["results"] if r["sample_id"] in ids]
            return f"{sum(r[key] in ok for r in rs)}/{len(rs)}" if rs else "–"

        ws.append(["Chuyển JSON: mẫu không có lỗi", d] + [pct(val[m]) for m in models])
        ws.append(["Dịch: mẫu qua kiểm tra tự động", d] + [pct(chk[m]) for m in models])
        if judge:
            jr = [r for r in judge["results"] if r["sample_id"] in ids]
            wins = collections.Counter(f["better_model"] for r in jr for f in r["result"]["fields"])
            nf = sum(wins.values()) or 1
            ws.append(["Giám khảo: % trường thắng (hòa: " + f"{wins['tie']/nf:.0%})", d] +
                      [f"{wins[m]/nf:.0%}" for m in models])
            avg = {m: [r["result"]["overall"]["score_by_model"][m] for r in jr] for m in models}
            ws.append(["Giám khảo: điểm TB (1–5)", d] +
                      [f"{sum(v)/len(v):.2f}" if v else "–" for v in avg.values()])
            sev = {m: collections.Counter(e["severity"] for r in jr for f in r["result"]["fields"]
                                          for e in f["errors"] if e["model"] == m) for m in models}
            ws.append(["Giám khảo: lỗi nghiêm trọng / nặng / nhẹ", d] +
                      [f"{sev[m]['nghiem_trong']} / {sev[m]['nang']} / {sev[m]['nhe']}" for m in models])
            old = [r["result"]["overall"]["score_by_model"]["old_ko"] for r in jr]
            if old and d == "(tất cả)":
                ws.append(["Giám khảo: điểm TB bản tiếng Hàn cũ", d, f"{sum(old)/len(old):.2f}"])
        if ed_scores and d == "(tất cả)":
            ws.append(["Biên tập viên: điểm TB (1–5)", d] +
                      [f"{sum(ed_scores[m])/len(ed_scores[m]):.2f}" if ed_scores.get(m) else "–" for m in models])
    style(ws, [44, 22] + [14] * len(models))

    # --- Chuyển JSON ---
    ws = wb.create_sheet("Chuyen_JSON")
    ws.append(["sample_id", "Dạng bài", "Model", "Kết quả", "Độ phủ tiếng Việt", "Lỗi", "Cảnh báo"])
    for m in models:
        for r in (val[m] or {"results": []})["results"]:
            ws.append([r["sample_id"], r.get("dang_bai", ""), m, r["status"],
                       f"{r['vi_coverage']:.1%}" if r.get("vi_coverage") is not None else "",
                       "\n".join(r["issues"]), "\n".join(r["warnings"])])
    style(ws, [24, 20, 9, 11, 12, 60, 60])

    # --- Chi tiết dịch ---
    ws = wb.create_sheet("Chi_tiet_dich")
    judged = {}
    if judge:
        for r in judge["results"]:
            for f in r["result"]["fields"]:
                judged[(r["sample_id"], f["id"])] = f
    ws.append(["sample_id", "Trường", "Tiếng Việt"] + [f"{m}" for m in models] +
              [f"Lỗi tự động {m}" for m in models] + ["Giám khảo chọn", "Ghi chú giám khảo"])
    fields_by = {m: {} for m in models}
    for m in models:
        for r in (chk[m] or {"results": []})["results"]:
            for f in r["fields"]:
                fields_by[m][(r["sample_id"], f["path"])] = f
    keys = sorted({k for m in models for k in fields_by[m]})
    for k in keys:
        f0 = next(fields_by[m][k] for m in models if k in fields_by[m])
        j = judged.get(k, {})
        notes = "\n".join(f"[{e['model']}/{e['severity']}/{e['type']}] {e['note']}" for e in j.get("errors", []))
        ws.append([k[0], k[1], f0["vi"]] + [fields_by[m].get(k, {}).get(lang, "") for m in models] +
                  ["; ".join(e["loai"] for e in fields_by[m].get(k, {}).get("errors", [])) for m in models] +
                  [j.get("better_model", ""), notes])
    style(ws, [22, 26, 50] + [50] * len(models) + [22] * len(models) + [12, 50])

    # --- Vấn đề bản gốc & ghi chú biên tập viên ---
    ws = wb.create_sheet("Van_de_ban_goc")
    ws.append(["sample_id", "Nguồn", "Nội dung"])
    if judge:
        for r in judge["results"]:
            if r["result"].get("van_de_ban_goc", "").strip():
                ws.append([r["sample_id"], "giám khảo", r["result"]["van_de_ban_goc"]])
    for sid, path_, note in ed_notes:
        ws.append([sid, f"biên tập viên ({path_})", note])
    style(ws, [24, 26, 90])

    out = p("reports", "bao_cao_so_sanh.xlsx")
    wb.save(out)
    print("→", out)

    # --- Phiếu chấm mù cho biên tập viên (chỉ tạo nếu chưa có, tránh ghi đè điểm đã chấm) ---
    sheet_path = p("reports", "phieu_bien_tap_vien.xlsx")
    if not os.path.exists(sheet_path) and all(chk[m] for m in models):
        rng = random.Random(cfg["seed"] + 1)
        ew = Workbook()
        cs = ew.active
        cs.title = "Cham_diem"
        cs.append(["sample_id", "Trường", "Tiếng Việt", "Bản A", "Bản B", "Điểm A (1–5)", "Điểm B (1–5)", "Ghi chú"])
        answer = []
        for k in keys:
            order = models[:]
            rng.shuffle(order)
            cs.append([k[0], k[1], fields_by[order[0]].get(k, {}).get("vi", "")] +
                      [fields_by[m].get(k, {}).get(lang, "") for m in order] + ["", "", ""])
            answer.append({"sample_id": k[0], "path": k[1], "A": order[0], "B": order[1]})
        style(cs, [22, 26, 50, 50, 50, 12, 12, 40])
        ew.save(sheet_path)
        save_json(answer, p("reports", "_dap_an_phieu_bien_tap_vien.json"))
        print("→", sheet_path, "(gửi file này cho biên tập viên)")
        print("→ reports/_dap_an_phieu_bien_tap_vien.json (đáp án A/B – KHÔNG gửi cho biên tập viên)")


if __name__ == "__main__":
    main()
